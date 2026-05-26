from pptx import Presentation
prs = Presentation('Naresh & Co Profile - Copy Final.pptx')
print('slides', len(prs.slides))
for idx, slide in enumerate(prs.slides, 1):
    print('--- Slide', idx, '---')
    for shape in slide.shapes:
        if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
            txt = shape.text.strip()
            if txt:
                print('TEXT:', repr(txt))
        if shape.shape_type == 13:
            print('PICTURE', shape.name, 'left', shape.left, 'top', shape.top, 'w', shape.width, 'h', shape.height)
    print()

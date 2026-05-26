from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
prs = Presentation('Naresh & Co Profile - Copy.pptx')
slide = prs.slides[3]
print('Slide 4 shapes:')
for shape in slide.shapes:
    print(shape.shape_id, shape.name, shape.shape_type, getattr(shape, 'left', None), getattr(shape, 'top', None), getattr(shape, 'width', None), getattr(shape, 'height', None))
    if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
        print('   text:', repr(shape.text))

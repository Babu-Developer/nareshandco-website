from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
prs = Presentation('Naresh & Co Profile - Copy.pptx')
print('slides', len(prs.slides))
for idx, slide in enumerate(prs.slides, 1):
    print('==== Slide', idx, '====')
    for shape in slide.shapes:
        print('shape id', shape.shape_id, 'type', shape.shape_type, 'name', shape.name)
        if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
            print('TEXT:')
            for para in shape.text_frame.paragraphs:
                print('  ', repr(para.text))
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print('PICTURE:', shape.name)
    print()

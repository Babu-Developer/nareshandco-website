import sys
import csv
from pptx import Presentation


def extract_shapes(prs):
    all_shapes = []
    for si, slide in enumerate(prs.slides, start=1):
        for idx, shape in enumerate(slide.shapes, start=1):
            try:
                left = int(shape.left)
                top = int(shape.top)
                width = int(shape.width)
                height = int(shape.height)
            except Exception:
                left = top = width = height = None
            text = ''
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                try:
                    text = ''.join(p.text for p in shape.text_frame.paragraphs)
                except Exception:
                    text = ''
            all_shapes.append({
                'slide': si,
                'index': idx,
                'type': shape.shape_type,
                'left': left,
                'top': top,
                'width': width,
                'height': height,
                'text': text.replace('\n', ' '),
            })
    return all_shapes


def find_by_key(shapes, slide, index):
    for s in shapes:
        if s['slide'] == slide and s['index'] == index:
            return s
    return None


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('Usage: generate_report.py original.pptx modified.pptx')
        sys.exit(2)
    orig = sys.argv[1]
    mod = sys.argv[2]
    out = 'pptx_change_report.csv'

    prs_orig = Presentation(orig)
    prs_mod = Presentation(mod)

    s_orig = extract_shapes(prs_orig)
    s_mod = extract_shapes(prs_mod)

    max_slide = max((s['slide'] for s in s_orig+s_mod), default=0)

    rows = []
    # match by slide and index where possible; also report unmatched shapes
    indices = set(( (s['slide'], s['index']) for s in s_orig+s_mod ))
    for slide, index in sorted(indices):
        o = find_by_key(s_orig, slide, index)
        m = find_by_key(s_mod, slide, index)
        row = {
            'slide': slide,
            'index': index,
            'left_orig': o['left'] if o else '',
            'top_orig': o['top'] if o else '',
            'width_orig': o['width'] if o else '',
            'height_orig': o['height'] if o else '',
            'text_orig': o['text'] if o else '',
            'left_mod': m['left'] if m else '',
            'top_mod': m['top'] if m else '',
            'width_mod': m['width'] if m else '',
            'height_mod': m['height'] if m else '',
            'text_mod': m['text'] if m else '',
        }
        # compute deltas if numeric
        def d(a,b):
            try:
                return (int(b) - int(a))
            except Exception:
                return ''
        row['delta_left'] = d(row['left_orig'], row['left_mod'])
        row['delta_top'] = d(row['top_orig'], row['top_mod'])
        row['delta_width'] = d(row['width_orig'], row['width_mod'])
        row['delta_height'] = d(row['height_orig'], row['height_mod'])
        rows.append(row)

    with open(out, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else [])
        if rows:
            writer.writeheader()
            for r in rows:
                writer.writerow(r)

    print('Report saved to', out)

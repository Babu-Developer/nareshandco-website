import sys
import os
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE


def fix_presentation(in_path, out_path):
    prs = Presentation(in_path)
    slide_w = prs.slide_width
    slide_h = prs.slide_height

    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    shape.left = int((slide_w - shape.width) / 2)
                    if shape.top < 0:
                        shape.top = Pt(20)

                if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                    tf = shape.text_frame
                    for p in tf.paragraphs:
                        p.alignment = PP_ALIGN.CENTER
                        for run in p.runs:
                            run.font.name = "Arial"
                            text_len = len(run.text or "")
                            if text_len <= 40:
                                run.font.size = Pt(28)
                            elif text_len <= 120:
                                run.font.size = Pt(20)
                            else:
                                run.font.size = Pt(16)
                    # center the textbox
                    shape.left = int((slide_w - shape.width) / 2)
            except Exception:
                # best-effort: ignore shapes we can't modify
                continue

    prs.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python fix_pptx.py <input.pptx> [output.pptx]")
        sys.exit(2)
    in_path = sys.argv[1]
    if not os.path.isfile(in_path):
        print(f"Input file not found: {in_path}")
        sys.exit(2)
    if len(sys.argv) >= 3:
        out_path = sys.argv[2]
    else:
        base = os.path.basename(in_path)
        out_path = os.path.join(os.getcwd(), "fixed_" + base)
    print(f"Processing: {in_path}\nSaving to: {out_path}")
    try:
        fix_presentation(in_path, out_path)
        print("Done.")
    except Exception as e:
        print("Failed:", e)
        sys.exit(1)
